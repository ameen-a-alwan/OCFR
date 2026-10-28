"""Document parsers for PDF, DOCX, and TXT.

SQLite is not an upload format — the app owns an internal SQLite DB for
registry of uploads and typed text entries (see ``app.core.database``).
"""

from __future__ import annotations

import logging
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable

import fitz  # PyMuPDF
from docx import Document

from app.core.config import Settings, get_settings
from app.services.ocr import OCRService, get_ocr_service

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}


@dataclass
class ParsedPage:
    """Extracted text for a logical page / section."""

    page_number: int
    text: str
    used_ocr: bool = False


@dataclass
class ParsedDocument:
    """Full parse result for a single source file."""

    filename: str
    file_type: str
    source_path: str
    pages: list[ParsedPage] = field(default_factory=list)

    @property
    def full_text(self) -> str:
        return "\n\n".join(p.text for p in self.pages if p.text.strip())


class DocumentParser:
    """Route files to the appropriate local parser."""

    def __init__(
        self,
        settings: Settings | None = None,
        ocr: OCRService | None = None,
    ) -> None:
        self.settings = settings or get_settings()
        self.ocr = ocr or get_ocr_service()

    def parse(self, path: Path) -> ParsedDocument:
        path = Path(path)
        if not path.exists():
            raise FileNotFoundError(f"File not found: {path}")

        suffix = path.suffix.lower()
        if suffix not in SUPPORTED_EXTENSIONS:
            raise ValueError(
                f"Unsupported file type '{suffix}'. "
                f"Supported uploads: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
            )

        if suffix == ".pdf":
            pages = self._parse_pdf(path)
            file_type = "pdf"
        elif suffix == ".docx":
            pages = self._parse_docx(path)
            file_type = "docx"
        else:
            pages = self._parse_txt(path)
            file_type = "txt"

        return ParsedDocument(
            filename=path.name,
            file_type=file_type,
            source_path=str(path.resolve()),
            pages=pages,
        )

    def _parse_pdf(self, path: Path) -> list[ParsedPage]:
        pages: list[ParsedPage] = []
        with fitz.open(path) as doc:
            for i, page in enumerate(doc, start=1):
                text = (page.get_text("text") or "").strip()
                used_ocr = False

                if len(text) < self.settings.pdf_text_min_chars:
                    try:
                        pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                        import numpy as np
                        from PIL import Image

                        image = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
                        ocr_text = self.ocr.extract_text_from_image(np.asarray(image))
                        if ocr_text.strip():
                            text = ocr_text.strip()
                            used_ocr = True
                    except Exception as exc:  # noqa: BLE001
                        logger.warning("OCR failed on %s page %d: %s", path.name, i, exc)

                if text:
                    pages.append(ParsedPage(page_number=i, text=text, used_ocr=used_ocr))
        return pages

    def _parse_docx(self, path: Path) -> list[ParsedPage]:
        document = Document(path)
        paragraphs = [p.text.strip() for p in document.paragraphs if p.text and p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
                if cells:
                    paragraphs.append(" | ".join(cells))

        text = "\n".join(paragraphs).strip()
        if not text:
            return []
        return [ParsedPage(page_number=1, text=text, used_ocr=False)]

    def _parse_txt(self, path: Path) -> list[ParsedPage]:
        raw = path.read_bytes()
        for encoding in ("utf-8", "utf-8-sig", "cp1256", "latin-1"):
            try:
                text = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:
            text = raw.decode("utf-8", errors="replace")

        text = text.strip()
        if not text:
            return []
        return [ParsedPage(page_number=1, text=text, used_ocr=False)]


def iter_supported_files(directory: Path) -> Iterable[Path]:
    directory = Path(directory)
    for path in sorted(directory.rglob("*")):
        if path.is_file() and path.suffix.lower() in SUPPORTED_EXTENSIONS:
            yield path
